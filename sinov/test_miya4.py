"""4-bosqich: xotira arxivi, ish usullari, signallar, ertalabki/haftalik xulosa, topshirish varaqasi, Excel."""
import asyncio
import io
import json
import os
import sys
from datetime import timedelta
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from miya import (ai, arxiv, baza, eksport, hisob, ishlar, maosh, odamlar, signallar, topshirish,  # noqa: E402
                  usullar, vadalar, vaqt, xulosa)


@pytest.fixture()
def ish(tmp_path):
    vaqt._siljish = timedelta(0)
    baza.boshlash(str(tmp_path / "miya.db"))
    b = odamlar.odam_yarat("Rustam Aliyev", "boshliq")
    k = odamlar.odam_yarat("Dilnoza Qodirova", "kassir", kim_id=b, maosh_tolaydi=1)
    a = odamlar.odam_yarat("Akmal Toshpo'latov", "xodim", kim_id=b)
    d = odamlar.odam_yarat("Dilshod Rahimov", "xodim", kim_id=b)
    kassa = hisob.kassa_yarat("Asosiy kassa", "naqd", "UZS", k, b)
    maosh.boshlangan_oy()
    hisob.tasdiqla(hisob.taklif("boshlangich_kassa", {"summa": 5_000_000, "kassa_id": kassa}, b), b)
    mid = baza.bajar("INSERT INTO mijoz(ism,yaratilgan) VALUES('Anvar Karimov',?)", vaqt.iso())
    yield {"b": b, "k": k, "a": a, "d": d, "kassa": kassa, "m": mid, "tmp": tmp_path}
    vaqt._siljish = timedelta(0)
    ai.KLIENT = None


def odam(i):
    return odamlar.odam(i)


def faol(hid, kim):
    if hisob.tasdiqla(hid, kim)["holat"] == "takror":
        hisob.faollashtir(hid, kim)
    return hid


def xabar(odam_id, matn_=None, stt=None, sozlar=None, tur="matn", fayl=None, vaqt_=None):
    return baza.bajar("INSERT INTO xabar(odam_id,tur,matn,stt_matn,stt_sozlar,fayl,davomiylik,vaqt,holat) "
                      "VALUES(?,?,?,?,?,?,?,?,'ishlandi')", odam_id, tur, matn_, stt,
                      json.dumps(sozlar) if sozlar else None, fayl, 20 if tur == "ovoz" else None, vaqt_ or vaqt.iso())


# ---------- X6: asl xabarni topish ----------

def test_ovozdan_topadi_va_soniyasi(ish):
    gap = "salom bugun do'konga bordim Karimov bilan narxni metri 45 mingdan kelishdik ertaga olib ketadi"
    sozlar = [{"soz": s, "dan": i * 0.8} for i, s in enumerate(gap.split())]
    vid = xabar(ish["a"], stt=gap, sozlar=sozlar, tur="ovoz", fayl="/yoq/ovoz.ogg")
    xabar(ish["a"], "Karimov 3 mln to'ladi")
    xabar(ish["d"], "Benzinga 300 ming ketdi")
    natija = arxiv.qidir("Karimov bilan narxni qanday kelishgan edik?", odam(ish["b"]))
    assert natija[0]["id"] == vid and natija[0]["ball"] == 2  # «qanday kelishgan edik» — savol so'zlari
    assert natija[0]["soniya"] == pytest.approx(4 * 0.8 - 1.0)  # «Karimov» 5-so'z — bir soniya oldindan
    assert "45 mingdan" in natija[0]["parcha"]


def test_kirill_va_qoshimcha(ish):
    xid = xabar(ish["a"], "Каримовга шартнома нусхасини бердим")
    assert arxiv.qidir("Karimov shartnoma", odam(ish["b"]))[0]["id"] == xid


def test_arxiv_huquqi(ish):
    xabar(ish["a"], "Akmalga 500 ming avans berdim")  # Akmal yozgan
    x2 = xabar(ish["d"], "Dilshodga avans kerak")
    assert [x["id"] for x in arxiv.qidir("avans", odam(ish["d"]))] == [x2]  # faqat o'ziniki (D-d)
    assert len(arxiv.qidir("avans", odam(ish["b"]))) == 2
    baza.bajar("UPDATE xabar SET yashirin=1 WHERE id=?", x2)
    assert arxiv.qidir("avans", odam(ish["d"])) == []


def test_xabar_qidir_asbobi(ish):
    xabar(ish["a"], "Karimov bilan narxni 45 mingdan kelishdik")
    d = ai.asbob_bajar(odam(ish["b"]), "xabar_qidir", {"sozlar": "Karimov narx"})
    assert d["xabarlar"][0]["kim"] == "Akmal Toshpo'latov"
    assert ai.asbob_bajar(odam(ish["d"]), "xabar_qidir", {"sozlar": "Karimov narx"})["xabarlar"] == []


# ---------- O5: ish usullari ----------

def test_ish_usuli_topiladi(ish):
    usullar.qosh("Kafel yotqizish", "Avval pol tekislanadi, keyin elim 2 sm qatlamda surtiladi, chok 3 mm.", ish["b"])
    usullar.qosh("Kassa yopish", "Kechqurun kassir pulni sanaydi va botga yozadi.", ish["b"])
    t = usullar.qidir("Kafelni qanday yotqizamiz?")
    assert t[0]["sarlavha"] == "Kafel yotqizish"
    d = ai.asbob_bajar(odam(ish["a"]), "ish_usuli", {"savol": "kafel qanday yotqiziladi"})
    assert "elim" in d["usullar"][0]["matn"]
    assert ai.asbob_bajar(odam(ish["a"]), "ish_usuli", {"savol": "Konditsioner tozalash"})["usullar"] == []
    with pytest.raises(usullar.UsulXato):
        usullar.qosh("", "  ", ish["b"])


# ---------- S7: yangi signallar ----------

def test_odatdan_katta_xarajat(ish):
    for n in range(6, 0, -1):
        vaqt._siljish = timedelta(days=-n * 5)
        faol(hisob.taklif("chiqim", {"summa": 300_000 + n * 1000, "kassa_id": ish["kassa"], "xarajat_turi": "Benzin"},
                          ish["b"]), ish["b"])
    vaqt._siljish = timedelta(0)
    hid = faol(hisob.taklif("chiqim", {"summa": 1_500_000, "kassa_id": ish["kassa"], "xarajat_turi": "Benzin"}, ish["b"]),
               ish["b"])
    s = signallar.katta_xarajatlar()
    assert [x["harakat_id"] for x in s] == [hid]
    assert "Benzin" in s[0]["matn"]


def test_takror_chek(ish):
    idlar = []
    for _ in range(2):
        xid = xabar(ish["k"], tur="rasm")
        hid = hisob.taklif("chiqim", {"summa": 120_000, "kassa_id": ish["kassa"], "xarajat_turi": "Ofis"}, ish["k"], xid)
        baza.bajar("INSERT INTO ilova(harakat_id,xabar_id,fayl,tur,xesh,vaqt) VALUES(?,?,?,?,?,?)", hid, xid, "x.jpg",
                   "rasm", "AQADbbb", vaqt.iso())
        hisob.faollashtir(hid, ish["k"])  # odam «Yo'q, boshqa» degan
        idlar.append(hid)
    s = signallar.takror_cheklar()
    assert len(s) == 1 and s[0]["harakat_id"] == idlar[1]


def test_band_va_bosh_xodim(ish):
    for i in range(6):
        ishlar.ish_yarat(f"Ish {i}", [(ish["a"], 1)], ish["b"])
    eski, _ = ishlar.ish_yarat("Eski ish", [(ish["d"], 1)], ish["b"])
    baza.bajar("UPDATE ish SET holat='qabul', bajarildi_vaqt=? WHERE id=?", vaqt.iso(vaqt.hozir() - timedelta(days=10)), eski)
    baza.bajar("INSERT INTO davomat(odam_id,sana,keldi) VALUES(?,?,?)", ish["d"], vaqt.bugun().isoformat(), vaqt.iso())
    s = {x["tur"]: x for x in signallar.band_boshlar()}
    assert s["band"]["odam_id"] == ish["a"] and "6 ta" in s["band"]["matn"]
    if not ishlar.dam_olishmi(vaqt.bugun()):
        assert s["bosh"]["odam_id"] == ish["d"]


def test_signallar_huquq_boyicha(ish):
    ishlar.ish_yarat("Kechikkan", [(ish["a"], 1)], ish["b"], muddat=(vaqt.bugun() - timedelta(days=2)).isoformat())
    assert any(x["tur"] == "ish_kech" for x in signallar.royxat(odam(ish["b"])))
    assert not any(x["tur"] == "ish_kech" for x in signallar.royxat(odam(ish["k"])))  # kassirga ish signali yo'q
    assert signallar.royxat(odam(ish["a"])) == []


# ---------- S6, S8: xulosa ----------

class SoxtaXulosaAI:
    def __init__(self, matn_):
        self.matn = matn_
        self.messages = self

    async def create(self, **kw):
        u = SimpleNamespace(input_tokens=10, output_tokens=10, cache_read_input_tokens=0, cache_creation_input_tokens=0)
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=self.matn)], usage=u, stop_reason="end_turn")


def test_ertalabki_xulosa_qolip_va_raqam_tekshiruvi(ish):
    vadalar.yoz(ish["m"], vaqt.bugun().isoformat(), ish["b"], summa=2_000_000)
    x = asyncio.run(xulosa.yasash("kun"))
    assert x["qolip"] == 1 and "5 000 000" in x["matn"].replace("\xa0", " ")
    assert "Anvar Karimov 2 000 000" in x["matn"].replace("\xa0", " ")
    ai.KLIENT = SoxtaXulosaAI("Xayrli tong!\nKassalarda 7 000 000 so'm.")  # noto'g'ri raqam
    x = asyncio.run(xulosa.yasash("kun", majburan=True))
    assert x["qolip"] == 1
    ai.KLIENT = SoxtaXulosaAI("Xayrli tong!\nKassalarda 5 000 000 so'm.\nBugun Karimov 2 mln to'lashi kerak.")
    x = asyncio.run(xulosa.yasash("kun", majburan=True))
    assert x["qolip"] == 0 and "Karimov 2 mln" in x["matn"]


def test_haftalik_xulosa_va_aniqlanmaganlar(ish):
    bugun = vaqt.bugun()
    dushanba = bugun - timedelta(days=bugun.weekday())
    vaqt._siljish = timedelta(days=-(bugun - (dushanba - timedelta(days=3))).days)  # o'tgan hafta payshanbasi
    faol(hisob.taklif("chiqim", {"summa": 400_000, "kassa_id": ish["kassa"], "xarajat_turi": "Ijara"}, ish["b"]), ish["b"])
    hid = hisob.taklif("avans", {"summa": 100_000, "kassa_id": ish["kassa"], "odam_id": ish["a"]}, ish["k"])
    hisob.tasdiqla(hid, ish["k"])  # Akmal hali «oldim» demagan
    baza.bajar("INSERT INTO karta(odam_id,chat_id,malumot,yaratilgan) VALUES(?,?,?,?)", ish["a"], 1, '{"tur":"kirim"}',
               vaqt.iso())
    vaqt._siljish = timedelta(0)
    f = xulosa.hafta_faktlar()
    assert f["hafta"]["xarajat"]["summa"] == 400_000 and f["oldingi_haftadan_farq"]["xarajat"]["summa"] == 400_000
    a = f["aniqlanmaganlar"]
    assert a["tasdiqlanmagan_soni"] == 2 and a["ochiq_kartalar_soni"] == 1  # kassir «berdim», Akmal «oldim»
    matn_ = xulosa.hafta_qolip(f).replace("\xa0", " ")
    assert "xarajat 400 000 so'm (+400 000 so'm)" in matn_ and "1 ta kartochka yopilmay qolgan" in matn_
    vaqt._siljish = timedelta(days=8)
    assert xulosa.kartalarni_eskirt() == 1


# ---------- O4: topshirish ----------

def test_topshirish_varaqasi_va_yopish(ish):
    kassa2 = hisob.kassa_yarat("Akmal kassasi", "naqd", "UZS", ish["a"], ish["b"])
    iid, _ = ishlar.ish_yarat("Eshik o'rnatish", [(ish["a"], 1)], ish["b"])
    baza.bajar("UPDATE mijoz SET masul_id=? WHERE id=?", ish["a"], ish["m"])
    vid = vadalar.yoz(ish["m"], vaqt.bugun().isoformat(), ish["b"], summa=1_000_000)
    faol(hisob.taklif("qolga_berish", {"summa": 200_000, "kassa_id": ish["kassa"], "odam_id": ish["a"]}, ish["a"]), ish["a"])
    hid = hisob.taklif("kirim", {"summa": 50_000, "kassa_id": kassa2}, ish["b"])
    hisob.tasdiqla(hid, ish["b"])  # Akmal kassasiga tushdi — «oldim» kutilmoqda
    odamlar.holat_qoy(ish["a"], "ketmoqda", ish["b"])
    v = topshirish.varaq(ish["a"])
    assert not v["tayyor"] and len(v["ishlar"]) == 1 and v["mijozlar"] and v["vadalar"] and v["kassalar"]
    assert v["qolida"][0]["summa"] == 200_000 and v["kutmoqda"]
    with pytest.raises(topshirish.TopshirishXato):
        topshirish.yopish(ish["a"], ish["b"])
    n = topshirish.otkaz(ish["a"], ish["d"], ish["b"])
    assert n == {"ishlar": 1, "mijozlar": 1, "vadalar": 1, "kassalar": 1}
    assert [x["odam_id"] for x in ishlar.ish(iid)["ijrochilar"]] == [ish["d"]]
    assert vadalar.vada(vid)["masul_id"] == ish["d"]
    assert hisob.harakat(hid)["oluvchi_id"] == ish["d"]  # tasdiq yangi mas'ulga o'tdi
    v = topshirish.varaq(ish["a"])
    assert not v["tayyor"] and not v["ishlar"] and v["qolida"]  # qo'lidagi pul hali topshirilmagan
    topshirish.yopish(ish["a"], ish["b"], "pulni naqd qaytardi, keyin yoziladi")
    assert odam(ish["a"])["holat"] == "ketgan"
    assert baza.bir("SELECT 1 FROM jurnal WHERE amal='topshirishsiz_yopildi'")


# ---------- E5: Excel ----------

def test_excel_son_va_formula_himoyasi(ish):
    faol(hisob.taklif("chiqim", {"summa": 250_000, "kassa_id": ish["kassa"], "xarajat_turi": "Ofis",
                                 "izoh": "=HYPERLINK(\"http://yomon\")"}, ish["k"]), ish["k"])
    wb = load_workbook(io.BytesIO(eksport.kassa(odam(ish["b"]), vaqt.bugun().isoformat(), vaqt.bugun().isoformat())))
    ws = wb["Yozuvlar"]
    qatorlar = list(ws.iter_rows(min_row=2, values_only=True))
    chiqim = [r for r in qatorlar if r[1] == "Xarajat"][0]
    assert chiqim[4] == 250_000 and isinstance(chiqim[4], int)
    izoh_katak = [c for c in ws[ws.max_row] if c.column == 12][0]
    assert izoh_katak.data_type == "s" and izoh_katak.value.startswith("=HYPERLINK")
    assert wb["Kassalar"]["D2"].value == 4_750_000
    sid, _ = hisob.sotuv_tuz({"summa": 900_000, "mijoz_id": ish["m"], "kassa_id": ish["kassa"]}, ish["b"])
    faol(sid, ish["b"])
    wb = load_workbook(io.BytesIO(eksport.qarzlar(odam(ish["b"]))))
    assert wb["Mijozlar qarzi"]["C2"].value == 900_000
    maosh.shart_qoy(ish["a"], "oylik", 4_000_000, "UZS", vaqt.oy_boshi(vaqt.oy(vaqt.bugun())).isoformat(), ish["b"])
    wb = load_workbook(io.BytesIO(eksport.maosh_(vaqt.oy(vaqt.bugun()))))
    assert wb["Varaqlar"]["A2"].value == "Akmal Toshpo'latov"
