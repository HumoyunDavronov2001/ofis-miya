"""Excel'ga chiqarish (E5, D-g): kassa, maosh, qarzlar.

Summalar son bo'lib yoziladi (so'm — butun, dollar — 2 xona). Foydalanuvchi matni hech qachon formula bo'lib qolmaydi.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import baza, hisob, maosh, mijozlar, odamlar, vadalar, vaqt

SARLAVHA_FON = PatternFill("solid", fgColor="EBF0FF")
HOLAT = {"faol": "Faol", "nizo": "Nizo", "bekor": "Bekor", "ruxsat_kutilmoqda": "Ruxsat kutmoqda", "taklif": "Taklif"}


def _pul(minor: int | None, valyuta: str):
    if minor is None:
        return None
    return round(minor / 100, 2) if valyuta == "USD" else int(minor)


def _varaq(wb: Workbook, nom: str, ustunlar: list[tuple[str, int]], birinchi: bool = False):
    ws = wb.active if birinchi else wb.create_sheet()
    ws.title = nom[:31]
    for i, (sarlavha, kenglik) in enumerate(ustunlar, 1):
        c = ws.cell(row=1, column=i, value=sarlavha)
        c.font = Font(bold=True)
        c.fill = SARLAVHA_FON
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = kenglik
    ws.freeze_panes = "A2"
    return ws


def _qator(ws, qiymatlar: list, valyutalar: dict[int, str] | None = None) -> None:
    r = ws.max_row + 1
    for i, v in enumerate(qiymatlar, 1):
        c = ws.cell(row=r, column=i)
        c.value = v
        if isinstance(v, str):
            c.data_type = "s"  # «=…» bilan boshlangan izoh formula emas — matn (D-g)
        elif isinstance(v, (int, float)) and valyutalar and i in valyutalar:
            c.number_format = "#,##0.00" if valyutalar[i] == "USD" else "#,##0"


def _bayt(wb: Workbook) -> bytes:
    f = io.BytesIO()
    wb.save(f)
    return f.getvalue()


def kassa(o: dict, dan: str, gacha: str) -> bytes:
    wb = Workbook()
    ws = _varaq(wb, "Yozuvlar", [("Sana", 12), ("Tur", 24), ("Qayerdan", 26), ("Qayerga", 26), ("Summa", 16),
                                 ("Valyuta", 9), ("Ikkinchi summa", 16), ("Valyuta 2", 9), ("Holat", 16), ("Tasdiq", 22),
                                 ("Yozgan", 20), ("Izoh", 40)], birinchi=True)
    for h in reversed(hisob.harakatlar(korish=o, dan=dan, gacha=gacha, limit=100000)):
        tasdiq = []
        for tomon in ("beruvchi", "oluvchi"):
            if h[f"{tomon}_holat"] == "kutilmoqda":
                tasdiq.append(f"{tomon} kutilmoqda")
            elif h[f"{tomon}_holat"] == "rad":
                tasdiq.append(f"{tomon} rad etdi")
        yozgan = odamlar.odam(h["yozgan_id"])
        _qator(ws, [h["voqea_sana"], hisob.SHABLON.get(h["tur"], (h["tur"],))[0], hisob.hisob_nomi(h["dan"]),
                    hisob.hisob_nomi(h["ga"]), _pul(h["summa"], h["valyuta"]), h["valyuta"],
                    _pul(h["summa2"], h["valyuta2"] or "UZS"), h["valyuta2"] or "", HOLAT.get(h["holat"], h["holat"]),
                    ", ".join(tasdiq), yozgan["ism"] if yozgan else "Tizim", h["izoh"] or ""], {5: h["valyuta"], 7: h["valyuta2"] or "UZS"})
    ws2 = _varaq(wb, "Kassalar", [("Kassa", 28), ("Turi", 12), ("Valyuta", 9), ("Qoldiq", 18), ("Mas'ul", 22)])
    for k in hisob.kassa_holati():
        if odamlar.hisob_korishi_mumkin(o, f"kassa:{k['id']}"):
            masul = odamlar.odam(k["masul_id"])
            _qator(ws2, [k["nom"], k["tur"], k["valyuta"], _pul(k["qoldiq"], k["valyuta"]), masul["ism"] if masul else ""],
                   {4: k["valyuta"]})
    return _bayt(wb)


def maosh_(oy: str) -> bytes:
    wb = Workbook()
    ws = _varaq(wb, "Varaqlar", [("Xodim", 26), ("O'tgan oydan", 16), ("Hisoblandi", 16), ("Oldi", 16),
                                 ("Tasdiq kutmoqda", 16), ("Qolgani", 16), ("Valyuta", 9)], birinchi=True)
    ws2 = _varaq(wb, "Qatorlar", [("Xodim", 26), ("Sana", 12), ("Nima", 40), ("Summa", 16), ("Valyuta", 9)])
    for o in odamlar.faol_odamlar() + baza.hammasi("SELECT * FROM odam WHERE holat='ketgan'"):
        if not maosh.shartlar(o["id"]) and not baza.qiymat("SELECT 1 FROM harakat WHERE dan=? OR ga=? LIMIT 1",
                                                           f"maosh:{o['id']}", f"maosh:{o['id']}"):
            continue
        v = maosh.varaq(o["id"], oy)
        val = v["valyuta"]
        _qator(ws, [o["ism"], _pul(v["otgan"], val), _pul(v["hisoblandi"], val), _pul(v["olindi"], val),
                    _pul(v["kutilmoqda_jami"], val), _pul(v["qolgani"], val), val], {i: val for i in range(2, 7)})
        for q in v["qatorlar"]:
            _qator(ws2, [o["ism"], q["sana"] or "", q["nom"], _pul(q["summa"], val), val], {4: val})
    return _bayt(wb)


def qarzlar(o: dict) -> bytes:
    wb = Workbook()
    ws = _varaq(wb, "Mijozlar qarzi", [("Mijoz", 28), ("Telefon", 18), ("Qarz", 16), ("Valyuta", 9), ("Muddat", 12),
                                       ("Va'dalari", 34), ("Mas'ul", 22)], birinchi=True)
    for q in hisob.mijoz_qarzlari("mijoz"):
        m = mijozlar.mijoz(q["id"])
        if not mijozlar.korishi_mumkin(o, m):
            continue
        masul = odamlar.odam(m.get("masul_id"))
        ish_ = vadalar.ishonch(q["id"])
        for val, s in q["qarz"].items():
            _qator(ws, [q["ism"], q["telefon"] or "", _pul(s, val), val, q["muddat"] or "", ish_["matn"],
                        masul["ism"] if masul else ""], {3: val})
    ws2 = _varaq(wb, "Bizning qarzimiz", [("Ta'minotchi", 28), ("Telefon", 18), ("Qarz", 16), ("Valyuta", 9), ("Muddat", 12)])
    if odamlar.boshliqmi(o) or odamlar.mumkin(o, "mijoz_qarz"):
        for q in hisob.mijoz_qarzlari("taminotchi"):
            for val, s in q["qarz"].items():
                _qator(ws2, [q["ism"], q["telefon"] or "", _pul(s, val), val, q["muddat"] or ""], {3: val})
    return _bayt(wb)


def fayl_nomi(tur: str, qism: str) -> str:
    return f"{tur}_{qism or vaqt.bugun().isoformat()}.xlsx"


def hisobot(d: dict) -> bytes:
    """CRM3: hisobot Excel'da — hisobot.davr() natijasidan (raqamlar o'sha yerda hisoblangan)."""
    wb = Workbook()
    f = d["foyda"]
    ws = _varaq(wb, "Foyda-zarar", [("Ko'rsatkich", 34), ("So'm", 18)], birinchi=True)
    for nom, k in (("Sotuv", "sotuv"), ("Xarajat", "xarajat"), ("Ish haqi (yozilgani)", "ishhaqi"), ("Foyda", "foyda"),
                   ("Kassaga tushdi", "kassa_kirim"), ("Kassadan chiqdi", "kassa_chiqim")):
        _qator(ws, [nom, f[k]], {2: "UZS"})
    _qator(ws, [f"Davr: {d['dan']} — {d['gacha']}", None])
    if f["oy_yopilmagan"]:
        _qator(ws, ["Diqqat: maoshi yopilmagan oy bor — ish haqi to'liq emas", None])
    ws = _varaq(wb, "Xarajatlar", [("Turi", 30), ("So'm", 18), ("Ulushi, %", 12)])
    for x in d["xarajat_turlar"]:
        _qator(ws, [x["nom"], x["summa"], x["ulush"]], {2: "UZS"})
    ws = _varaq(wb, "Sotuv — mijozlar", [("Mijoz", 30), ("Sotuvlar soni", 14), ("So'm", 18)])
    for x in d["sotuv_mijoz"]:
        _qator(ws, [x["ism"], x["soni"], x["summa"]], {3: "UZS"})
    ws = _varaq(wb, "Sotuv — mahsulotlar", [("Mahsulot / xizmat", 34), ("Miqdor", 12), ("Birlik", 10), ("So'm", 18)])
    for x in d.get("sotuv_mahsulot", []):
        _qator(ws, [x["nom"], x["miqdor"], x["birlik"], x["summa"]], {4: "UZS"})
    ws = _varaq(wb, "Sotuv — sotuvchilar", [("Sotuvchi", 30), ("Sotuvlar soni", 14), ("So'm", 18)])
    for x in d["sotuv_sotuvchi"]:
        _qator(ws, [x["ism"], x["soni"], x["summa"]], {3: "UZS"})
    ws = _varaq(wb, "Kassalar", [("Kassa", 26), ("Valyuta", 9), ("Boshida", 16), ("Kirim", 16), ("Chiqim", 16), ("Oxirida", 16)])
    for k in d["kassalar"]:
        v = k["valyuta"]
        _qator(ws, [k["nom"], v, _pul(k["boshida"], v), _pul(k["kirim"], v), _pul(k["chiqim"], v), _pul(k["oxirida"], v)],
               {3: v, 4: v, 5: v, 6: v})
    ws = _varaq(wb, "Qarzlar muddati", [("Guruh", 22), ("Mijoz", 30), ("So'm", 18), ("Muddat", 12)])
    for b in d["qarz_yoshi"]:
        for m in b["mijozlar"]:
            _qator(ws, [b["nom"], m["ism"], m["summa"], m["muddat"] or ""], {3: "UZS"})
    return _bayt(wb)
